"""
آگهی‌های خودروی دیوار تهران با قیمت مشخص و زیر سقف قیمت -> اکسل.

استفاده:
    python divar_cars.py                              # سقف پیش‌فرض: ۲۰٬۰۰۰٬۰۰۰٬۰۰۰ ریال
    python divar_cars.py --max-rial 15000000000 -o cars.xlsx
    python divar_cars.py --limit 50                   # ۵۰ آگهی به‌جای ۱۰ تای پیش‌فرض
    python divar_cars.py --limit 0                    # همه آگهی‌ها (تا --max-pages)
    python divar_cars.py --max-pages 5 --no-details   # سریع: فقط اطلاعات لیست

بدون مرورگر؛ فقط requests + BeautifulSoup + openpyxl.

- فقط شهر تهران و دسته «سواری و وانت» (light).
- آگهی‌های «توافقی» یا بدون قیمت حذف می‌شوند.
- قیمت‌های دیوار به تومان است؛ سقف به ریال گرفته و تبدیل می‌شود.
- برای هر آگهی صفحه جزئیاتش هم خوانده می‌شود (کارکرد، مدل، رنگ، گیربکس، بدنه،
  بیمه، توضیحات و ...) و هر فیلدی که پیدا شود یک ستون اکسل می‌شود.
"""

import argparse
import json
import re
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TEHRAN_CITY_ID = "1"
CATEGORY = "light"  # سواری و وانت
SEARCH_API = "https://api.divar.ir/v8/postlist/w/search"
POST_API = "https://api.divar.ir/v8/posts-v2/web/{token}"
POST_URL = "https://divar.ir/v/{token}"
LIST_PAGE = "https://divar.ir/s/tehran/car"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.8",
    "Origin": "https://divar.ir",
    "Referer": "https://divar.ir/",
}
TIMEOUT = 20
DELAY = 1.5  # مکث بین درخواست‌ها (ثانیه) تا دیوار بلاک نکند

PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

session = requests.Session()
session.headers.update(HEADERS)


def to_int(text):
    """'۸۵۰٬۰۰۰٬۰۰۰ تومان' -> 850000000 ؛ اگر عددی نبود None."""
    if not text:
        return None
    digits = re.sub(r"\D", "", str(text).translate(PERSIAN_DIGITS))
    return int(digits) if digits else None


def price_from_texts(*texts):
    """از بین متن‌های کارت آگهی، متنی که «تومان» دارد قیمت است؛ «توافقی» یعنی بدون قیمت."""
    for text in texts:
        if text and "تومان" in text:
            return to_int(text)
    return None


# ---------------------------------------------------------------- لیست آگهی‌ها
def search_body(max_toman, pagination_data=None):
    body = {
        "city_ids": [TEHRAN_CITY_ID],
        "source_view": "CATEGORY",
        "search_data": {
            "form_data": {
                "data": {
                    "category": {"str": {"value": CATEGORY}},
                    "price": {"number_range": {"maximum": str(max_toman)}},
                }
            }
        },
    }
    if pagination_data:
        body["pagination_data"] = pagination_data
    return body


def parse_post_rows(data):
    """ویجت‌های POST_ROW پاسخ API -> لیست آگهی."""
    posts = []
    for widget in data.get("list_widgets") or []:
        if widget.get("widget_type") != "POST_ROW":
            continue
        w = widget.get("data") or {}
        token = w.get("token") or ((w.get("action") or {}).get("payload") or {}).get("token")
        if not token:
            continue
        top = w.get("top_description_text") or ""
        middle = w.get("middle_description_text") or ""
        bottom = w.get("bottom_description_text") or ""
        posts.append({
            "token": token,
            "عنوان": w.get("title") or "",
            "قیمت (تومان)": price_from_texts(middle, top, bottom),
            "کارکرد (لیست)": top if "کیلومتر" in top else "",
            "زمان و محله": bottom,
            "لینک": POST_URL.format(token=token),
        })
    return posts


def qualifying(posts, max_toman):
    """فقط قیمت‌دار و زیر سقف (حتی اگر فیلتر سایت اعمال نشده باشد)؛ بدون آگهی تکراری."""
    unique = {}
    for post in posts:
        price = post["قیمت (تومان)"]
        if price and price <= max_toman and post["token"] not in unique:
            post["قیمت (ریال)"] = price * 10
            unique[post["token"]] = post
    return list(unique.values())


def fetch_list_api(max_toman, max_pages, limit):
    posts, pagination = [], None
    for page in range(1, max_pages + 1):
        resp = session.post(SEARCH_API, json=search_body(max_toman, pagination), timeout=TIMEOUT)
        resp.raise_for_status()
        data = resp.json()
        rows = parse_post_rows(data)
        print(f"صفحه {page}: {len(rows)} آگهی")
        if not rows:
            break
        posts += rows
        if limit and len(qualifying(posts, max_toman)) >= limit:
            break  # به تعداد خواسته‌شده رسیدیم؛ صفحه بعد لازم نیست
        pagination = data.get("pagination", {}).get("data")
        if not pagination or data.get("pagination", {}).get("has_next_page") is False:
            break
        time.sleep(DELAY)
    return posts


def fetch_list_html(max_toman):
    """
    اگر API جواب نداد: صفحه HTML دیوار (فقط صفحه اول، چون بقیه با جاوااسکریپت می‌آید).
    """
    resp = session.get(LIST_PAGE, params={"price": f"-{max_toman}"}, timeout=TIMEOUT)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    posts = []
    for card in soup.select("article.kt-post-card, div.kt-post-card"):
        link = card.find("a", href=re.compile(r"/v/"))
        if link is None:
            link = card.find_parent("a", href=re.compile(r"/v/"))
        if link is None:
            continue
        token = link["href"].rstrip("/").split("/")[-1]
        title_el = card.select_one(".kt-post-card__title, h2")
        texts = [d.get_text(" ", strip=True) for d in card.select(".kt-post-card__description, .kt-post-card__bottom-description")]
        posts.append({
            "token": token,
            "عنوان": title_el.get_text(strip=True) if title_el else "",
            "قیمت (تومان)": price_from_texts(*texts),
            "کارکرد (لیست)": next((t for t in texts if "کیلومتر" in t), ""),
            "زمان و محله": texts[-1] if texts else "",
            "لینک": POST_URL.format(token=token),
        })
    return posts


# ---------------------------------------------------------------- جزئیات آگهی
def walk_fields(node, fields, description, in_description=False):
    """
    هر جای JSON که جفت title/value بود یک فیلد است (کارکرد، مدل، رنگ، گیربکس، ...).
    متن توضیحات از ویجت DESCRIPTION_ROW یا هر متنی داخل بخش DESCRIPTION جمع می‌شود.
    """
    if isinstance(node, dict):
        in_description = in_description or node.get("section_name") == "DESCRIPTION"
        title, value = node.get("title"), node.get("value")
        if isinstance(title, str) and isinstance(value, str) and title.strip() and value.strip():
            fields.setdefault(title.strip(), value.strip())
        data = node.get("data") if isinstance(node.get("data"), dict) else {}
        text = data.get("text") or (node.get("text") if in_description else None)
        if isinstance(text, str) and text.strip() and (
            node.get("widget_type") == "DESCRIPTION_ROW" or in_description
        ):
            description.append(text.strip())
        for child in node.values():
            walk_fields(child, fields, description, in_description)
    elif isinstance(node, list):
        for child in node:
            walk_fields(child, fields, description, in_description)


def description_from_html(token):
    """اگر API توضیحات نداد: خود صفحه آگهی در divar.ir/v/... (متن توضیحات یا meta description)."""
    resp = session.get(POST_URL.format(token=token), timeout=TIMEOUT)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    el = soup.select_one(".kt-description-row__text, [class*=description-row__text]")
    if el and el.get_text(strip=True):
        return el.get_text("\n", strip=True)
    meta = soup.select_one('meta[name="description"], meta[property="og:description"]')
    return (meta.get("content") or "").strip() if meta else ""


def fetch_details(token):
    fields, description = {}, []
    try:
        resp = session.get(POST_API.format(token=token), timeout=TIMEOUT)
        resp.raise_for_status()
        data = resp.json()
        walk_fields(data.get("sections") or data, fields, description)
        if not description:
            seo_desc = (data.get("seo") or {}).get("description")
            if seo_desc:
                description.append(seo_desc)
    except (requests.RequestException, ValueError) as exc:
        print(f"[API جزئیات {token} جواب نداد: {exc}] — خواندن توضیحات از صفحه آگهی")
    if not description:
        try:
            html_desc = description_from_html(token)
            if html_desc:
                description.append(html_desc)
        except requests.RequestException as exc:
            print(f"[صفحه آگهی {token} باز نشد: {exc}]")
    fields["توضیحات"] = "\n".join(dict.fromkeys(description))  # بدون متن تکراری
    return fields


# ---------------------------------------------------------------- اکسل
BASE_COLUMNS = ["عنوان", "قیمت (تومان)", "قیمت (ریال)", "توضیحات", "زمان و محله"]
PREFERRED = ["برند و تیپ", "مدل (سال تولید)", "کارکرد", "رنگ", "نوع سوخت", "گیربکس",
             "وضعیت موتور", "وضعیت شاسی‌ها", "وضعیت بدنه", "مهلت بیمهٔ شخص ثالث",
             "قیمت پایه", "معاوضه"]


def save_excel(posts, path):
    extra = []
    for post in posts:
        for key in post:
            if key not in BASE_COLUMNS + ["token", "لینک", "کارکرد (لیست)"] and key not in extra:
                extra.append(key)
    extra.sort(key=lambda k: PREFERRED.index(k) if k in PREFERRED else len(PREFERRED))
    if not any(p.get("کارکرد") for p in posts):
        extra.insert(0, "کارکرد (لیست)")
    # لینک درست کنار عنوان، تا بدون اسکرول بشود هر ردیف را در سایت چک کرد
    columns = BASE_COLUMNS[:1] + ["لینک"] + BASE_COLUMNS[1:] + extra

    wb = Workbook()
    ws = wb.active
    ws.title = "خودرو تهران"
    ws.sheet_view.rightToLeft = True
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "C2"  # عنوان و لینک موقع اسکرول افقی ثابت می‌مانند

    for post in posts:
        ws.append([post.get(col, "") for col in columns])
        r = ws.max_row
        for col in ("قیمت (تومان)", "قیمت (ریال)"):
            ws.cell(row=r, column=columns.index(col) + 1).number_format = "#,##0"
        ws.cell(row=r, column=columns.index("توضیحات") + 1).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        for col in ("عنوان", "لینک"):  # خود عنوان هم قابل کلیک است
            cell = ws.cell(row=r, column=columns.index(col) + 1)
            cell.hyperlink = post["لینک"]
            cell.style = "Hyperlink"

    widths = {"عنوان": 35, "توضیحات": 60, "لینک": 30, "زمان و محله": 25}
    for idx, col in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = widths.get(col, 16)
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="خودروهای دیوار تهران زیر سقف قیمت -> اکسل")
    parser.add_argument("--max-rial", type=int, default=20_000_000_000, help="سقف قیمت به ریال")
    parser.add_argument("--limit", type=int, default=10, help="تعداد آگهی (پیش‌فرض ۱۰؛ ۰ یعنی همه)")
    parser.add_argument("--max-pages", type=int, default=50, help="حداکثر صفحات لیست (هر صفحه حدود ۲۰ آگهی)")
    parser.add_argument("--no-details", action="store_true", help="جزئیات آگهی‌ها خوانده نشود (سریع‌تر)")
    parser.add_argument("-o", "--output", default="divar_cars.xlsx", help="فایل اکسل خروجی")
    args = parser.parse_args()
    max_toman = args.max_rial // 10

    try:
        posts = fetch_list_api(max_toman, args.max_pages, args.limit)
    except (requests.RequestException, ValueError) as exc:
        print(f"[API دیوار جواب نداد: {exc}] — خواندن صفحه HTML (فقط صفحه اول)...")
        posts = fetch_list_html(max_toman)

    posts = qualifying(posts, max_toman)
    if args.limit:
        posts = posts[:args.limit]
    print(f"\n{len(posts)} آگهی قیمت‌دار زیر {args.max_rial:,} ریال")
    if args.no_details:
        for i, post in enumerate(posts, start=1):
            print(f"{i}/{len(posts)}  {post['عنوان']}  |  {post['قیمت (تومان)']:,} تومان  |  {post['لینک']}")

    if not args.no_details:
        for i, post in enumerate(posts, start=1):
            try:
                post.update({k: v for k, v in fetch_details(post["token"]).items() if k not in post})
            except (requests.RequestException, ValueError) as exc:
                print(f"[خطا در جزئیات {post['token']}: {exc}]")
            print(f"{i}/{len(posts)}  {post['عنوان']}  |  {post['قیمت (تومان)']:,} تومان  |  {post['لینک']}")
            time.sleep(DELAY)

    save_excel(posts, args.output)
    print(f"\nذخیره شد: {args.output}")


if __name__ == "__main__":
    main()
