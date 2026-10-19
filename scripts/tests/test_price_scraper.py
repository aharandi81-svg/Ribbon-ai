"""تست‌های واحد پارس قیمت/وزن/تطبیق نام در price_scraper (بدون شبکه).

اجرا:  python -m pytest scripts/tests
"""

import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import price_scraper as ps  # noqa: E402


# ---------------------------------------------------------------- to_int_price (۱)
@pytest.mark.parametrize("text, expected", [
    ("۱۲۵,۰۰۰ تومان", 125000),
    ("۲۰٪ ۱۲۵,۰۰۰ تومان", 125000),            # درصد تخفیف قاطی نشود
    ("-20% 125,000 تومان", 125000),
    ("۱۲۵,۰۰۰ ۱۰۰,۰۰۰ تومان", 125000),        # فقط اولین عدد
    ("۱۲۵٬۰۰۰٫۰۰", 125000),                    # اعشار دور ریخته شود
    ("125,000.00 تومان", 125000),
    ("١٢٥٬٠٠٠ تومان", 125000),                  # ارقام عربی
    ("125000", 125000),
    ("تخفیف ۲۰ درصد ۱۲۵,۰۰۰", 125000),        # «درصد» هم رد شود
    ("۹۹٫۰۰۰ تومان", 99000),                   # هزارگان با ممیز فارسی
    ("1.250.000 ریال", 1250000),               # هزارگان با نقطه
    ("12.50", 12),
    ("ناموجود", None),
    ("", None),
    (None, None),
])
def test_to_int_price(text, expected):
    assert ps.to_int_price(text) == expected


# ---------------------------------------------------------------- currency (۲)
@pytest.mark.parametrize("text, currency, expected", [
    ("۱,۲۵۰,۰۰۰ ریال", None, 125000),
    ("1,250,000 ﷼", None, 125000),
    ("1,250,000 IRR", None, 125000),
    ("۱۲۵,۰۰۰ تومان", None, 125000),
    ("1,250,000", "rial", 125000),             # واحد صریح بر متن مقدم است
    ("125,000", "toman", 125000),
])
def test_price_to_toman(text, currency, expected):
    assert ps.price_to_toman(text, currency) == expected


# ---------------------------------------------------------------- parse_weight (۳، ۴)
@pytest.mark.parametrize("title, expected", [
    ("شکر سفید ۹۰۰ گرمی", (900, "گرم")),
    ("روغن ۱/۵ لیتری", (1500, "میلی‌لیتر")),
    ("آرد ۱ کیلوگرمی", (1000, "گرم")),
    ("برنج ۱۰ کیلویی", (10000, "گرم")),
    ("زعفران ۵۰۰ میلی گرم", (0.5, "گرم")),
    ("۱٬۰۰۰ گرمی", (1000, "گرم")),            # جداکننده هزارگان
    ("شیر 1,500 میلی لیتر", (1500, "میلی‌لیتر")),
    ("بسته ۶ عددی ۴۰۰ گرمی", (400, "گرم")),    # بدون × ضرب نشود
    ("نوشابه ۳۰۰ میلی لیتر بسته ۶ عددی", (300, "میلی‌لیتر")),
    ("۶ عدد × ۴۰۰ گرم", (2400, "گرم")),        # الگوی صریح
    ("۴۰۰ گرمی × ۶ عدد", (2400, "گرم")),
    ("۲ کیلو و ۵۰۰ گرم", (2000, "گرم")),       # اولین واحد به ترتیب متن
    ("قند", (None, None)),
    ("", (None, None)),
])
def test_parse_weight(title, expected):
    assert ps.parse_weight(title) == expected


# ---------------------------------------------------------------- name matching (۵)
@pytest.mark.parametrize("query, title", [
    ("زردچوبه", "پودر زردچوبه ۹۰ گرمی"),
    ("دارچین", "دارچین پودر شده"),
    ("فلفل سیاه", "فلفل سیاه دانه"),
    ("سیر", "سیر پودر شده"),
    ("شکر", "شکر سفید ۹۰۰ گرمی"),
    ("روغن مایع", "روغن مایع لادن"),
    ("برنج ایرانی", "برنج ایرانی هاشمی"),
])
def test_spices_and_brands_are_comparable(query, title):
    best = ps.pick_best(query, [{"title": title, "price": 1, "url": "u"}])
    assert best is not None and best["comparable"] and best["match"].startswith("دقیق")


@pytest.mark.parametrize("query, title", [
    ("گوجه فرنگی", "سس گوجه فرنگی"),
    ("سیب زمینی", "چیپس سیب زمینی"),
    ("لیمو", "اسانس لیمو"),
    ("ایرانی برنج", "برنج ایرانی"),               # ترتیب کلمات فرق دارد
])
def test_negative_or_gap_is_approximate(query, title):
    best = ps.pick_best(query, [{"title": title, "price": 1, "url": "u"}])
    assert best["match"].startswith("تقریبی") and not best["comparable"]


def test_negative_word_ranks_lower_not_rejected():
    best = ps.pick_best("گوجه فرنگی", [
        {"title": "سس گوجه فرنگی", "price": 1, "url": "a"},
        {"title": "گوجه فرنگی گلخانه ای ممتاز", "price": 2, "url": "b"},
    ])
    assert best["url"] == "b"


def test_negative_word_in_query_is_allowed():
    best = ps.pick_best("سس مایونز", [{"title": "سس مایونز ۴۵۰ گرمی", "price": 1, "url": "u"}])
    assert best["match"] == "دقیق" and best["comparable"]


def test_pick_best_prefers_fewer_extras():
    best = ps.pick_best("کرفس", [
        {"title": "دانه کرفس ۱۰۰ گرمی", "price": 1, "url": "a"},
        {"title": "کرفس ۵۰۰ گرمی", "price": 2, "url": "b"},
    ])
    assert best["url"] == "b" and best["match"] == "دقیق"


def test_collect_per_unit(monkeypatch):
    monkeypatch.setattr(ps, "DELAY", 0)
    monkeypatch.setattr(ps, "SITES", {
        "brand": lambda q: [{"title": "شکر سفید ۹۰۰ گرمی", "price": 90000, "url": "u1"}],
        "approx": lambda q: [{"title": "سس شکر ۵۰۰ گرمی", "price": 50000, "url": "u2"}],
    })
    brand, approx = ps.collect(["شکر"])
    assert brand["per_unit"] == 100000              # تومان به ازای هر کیلو (۷)
    assert approx["match"].startswith("تقریبی") and approx["per_unit"] is None
    assert approx["price"] == 50000


# ---------------------------------------------------------------- parsehshop JSON (۱، ۲، ۶)
def test_parseh_json_uses_price_amount_and_currency():
    data = {"products": [
        {"name": "شکر ۹۰۰ گرمی", "price": "۹۰۰,۰۰۰ ریال", "price_amount": 900000.0, "url": "u1"},
        {"name": "آرد", "price": "۲۰٪ ۵۰,۰۰۰ تومان", "price_amount": 50000, "url": "u2"},
        {"name": "روغن", "price": "۱۲۵٬۰۰۰٫۰۰ تومان", "url": "u3"},   # بدون price_amount
    ]}
    assert [c["price"] for c in ps.parse_parseh_json(data)] == [90000, 50000, 125000]


def test_parseh_json_explicit_currency_field():
    data = {"currency": {"iso_code": "IRR", "sign": "﷼"},
            "products": [{"name": "شکر", "price": "900,000", "price_amount": 900000, "url": "u"}]}
    assert ps.parse_parseh_json(data)[0]["price"] == 90000


def test_parseh_json_skips_unavailable():
    data = {"products": [
        {"name": "a", "price_amount": 10000, "availability": "unavailable"},
        {"name": "b", "price_amount": 10000, "available_for_order": "0"},
        # مدیریت موجودی خاموش: quantity=0 ولی availability=available -> موجود
        {"name": "c", "price_amount": 10000, "availability": "available", "quantity": 0, "allow_oosp": 0},
        {"name": "d", "price_amount": 10000, "quantity": 0},
        {"name": "e", "price_amount": 10000, "availability": "last_remaining_items", "available_for_order": 0},
    ]}
    assert [c["title"] for c in ps.parse_parseh_json(data)] == ["c", "d", "e"]


def test_min_price_rejected():
    data = {"products": [{"name": "a", "price": "۹۹ تومان", "url": "u"},
                         {"name": "b", "price_amount": 500, "url": "u"}]}
    assert ps.parse_parseh_json(data) == []
    dk = {"data": {"products": [{"id": 1, "title_fa": "x", "default_variant": {"price": {"selling_price": 9000}}}]}}
    assert ps.parse_digikala_json(dk) == []


@pytest.mark.parametrize("data", [None, [], {}, {"products": None}])
def test_parseh_json_empty(data):
    assert ps.parse_parseh_json(data) == []


# ---------------------------------------------------------------- parsehshop HTML (۱، ۲، ۶)
HTML_17 = """
<script>var prestashop = {"currency":{"name":"Rial","iso_code":"IRR","sign":"ریال"}};</script>
<article class="product-miniature">
  <h3 class="product-title"><a href="/p/1-sugar.html">شکر ۹۰۰ گرمی</a></h3>
  <div class="product-price-and-shipping">
    <span class="regular-price">1,000,000</span>
    <span class="discount-percentage">-10%</span>
    <span class="price">900,000</span>
  </div>
</article>
<article class="product-miniature">
  <h3 class="product-title"><a href="/p/2-flour.html">آرد</a></h3>
  <ul class="product-flags"><li class="product-flag out_of_stock">ناموجود</li></ul>
  <div class="product-price-and-shipping"><span class="price">500,000</span></div>
</article>
"""


def test_parseh_html_discount_currency_and_stock():
    rows = ps.parse_parseh_html(HTML_17, "https://parsehshop.ir/search")
    assert rows == [{"title": "شکر ۹۰۰ گرمی", "price": 90000, "url": "https://parsehshop.ir/p/1-sugar.html"}]


def test_parseh_html_in_stock_without_flags():
    html = """<article class="product-miniature"><h3 class="product-title"><a href="/x">قند</a></h3>
    <span class="price">۱۲۵,۰۰۰ تومان</span></article>"""
    assert len(ps.parse_parseh_html(html)) == 1


def test_parseh_html_toman_text():
    html = """<article class="product-miniature"><h3 class="product-title"><a href="/x">قند</a></h3>
    <span class="price">۲۰٪ ۱۲۵,۰۰۰ تومان</span></article>"""
    assert ps.parse_parseh_html(html)[0]["price"] == 125000


# ---------------------------------------------------------------- digikala (۹)
@pytest.mark.parametrize("data", [{"data": None}, {}, None, {"data": {"products": None}}])
def test_digikala_empty_data(data):
    assert ps.parse_digikala_json(data) == []


def test_digikala_rial_to_toman():
    data = {"data": {"products": [
        {"id": 1, "title_fa": "شکر", "default_variant": {"price": {"selling_price": 1250000}},
         "url": {"uri": "/product/dkp-1/"}},
        {"id": 2, "title_fa": "ناموجود", "default_variant": []},
    ]}}
    assert ps.parse_digikala_json(data) == [
        {"title": "شکر", "price": 125000, "url": "https://www.digikala.com/product/dkp-1/"}
    ]


# ---------------------------------------------------------------- --debug (۵)
class FakeResp:
    def __init__(self, text, url="https://x/"):
        self.text, self.url, self.status_code = text, url, 200

    def raise_for_status(self):
        pass

    def json(self):
        import json
        return json.loads(self.text)


def test_debug_saves_raw_responses(tmp_path, monkeypatch):
    import json
    search = json.dumps({"products": [{"name": "شکر", "price_amount": 90000, "url": "https://parsehshop.ir/p/1"}]})
    dk = json.dumps({"data": {"products": [{"id": 7, "title_fa": "شکر",
                                            "default_variant": {"price": {"selling_price": 900000}}}]}})
    pages = {"https://parsehshop.ir/p/1": "<html>product</html>",
             "https://api.digikala.com/v1/search/": dk,
             "https://api.digikala.com/v2/product/7/": "{\"product\": 7}"}

    def fake_get(url, params=None, **kwargs):
        if url == ps.PARSEH_SEARCH:
            return FakeResp(search if (params or {}).get("ajax") else "<html>search</html>", url)
        return FakeResp(pages[url], url)

    monkeypatch.setattr(ps.session, "get", fake_get)
    monkeypatch.setattr(ps, "DEBUG_DIR", str(tmp_path))
    monkeypatch.setattr(ps, "DEBUG_SAVED", set())
    ps.search_parsehshop("شکر")
    ps.search_digikala("شکر")
    assert sorted(os.listdir(tmp_path)) == [
        "digikala_product.json", "digikala_search.json",
        "parseh_product.html", "parseh_search.html", "parseh_search.json",
    ]
    assert (tmp_path / "parseh_search.json").read_text(encoding="utf-8") == search


def test_debug_off_saves_nothing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(ps, "DEBUG_DIR", None)
    ps.debug_save("x.json", "{}")
    assert os.listdir(tmp_path) == []


# ---------------------------------------------------------------- bonakbazar / dastavval (عمومی)
def test_generic_json_rial_and_unit_per_kg():
    # شبیه بنک‌بازار: قیمت واحد ریالی به ازای هر کیلوگرم، وزن در فیلد unit
    data = {"data": {"items": [
        {"id": 408, "title": "برنج سوپر باسماتی هایلی", "unit_price": 550000, "unit": "کیلوگرم"},
        {"id": 9, "title": "روغن سرخ‌کردنی", "price": "7,150,000 ریال", "unit": {"title": "کارتن ۱۶ کیلویی"}},
        {"id": 3, "title": "بدون قیمت", "price": 0},
    ]}}
    found = ps.candidates_from_json(data, "rial", "https://bonakbazar.com/products/{id}")
    assert found[0] == {"title": "برنج سوپر باسماتی هایلی", "price": 55000,
                        "url": "https://bonakbazar.com/products/408", "weight_text": "1 کیلوگرم"}
    assert found[1]["price"] == 715000 and ps.parse_weight(found[1]["weight_text"]) == (16000, "گرم")
    assert len(found) == 2


def test_generic_json_toman_text_overrides_site_default():
    found = ps.candidates_from_json([{"name": "شکر ۹۰۰ گرمی", "price": "۹۰,۰۰۰ تومان", "url": "/p/1"}],
                                    "rial", None, "https://x.ir/search")
    assert found[0]["price"] == 90000 and found[0]["url"] == "https://x.ir/p/1"


def test_generic_html_json_ld_and_next_data():
    ld = ('<script type="application/ld+json">{"@type": "Product", "name": "زعفران ۴ گرمی", '
          '"url": "/p/2", "offers": {"price": "1200000", "priceCurrency": "IRR"}}</script>')
    found = ps.candidates_from_html(ld, "https://x.ir/")
    assert found == [{"title": "زعفران ۴ گرمی", "price": 120000, "url": "https://x.ir/p/2", "weight_text": ""}]
    nd = ('<script id="__NEXT_DATA__" type="application/json">'
          '{"props": {"products": [{"title": "عدس ۱ کیلویی", "price": 85000, "slug": "/product/ads"}]}}</script>')
    assert ps.candidates_from_html(nd, "https://x.ir/")[0]["price"] == 85000


def test_generic_html_cards_ignore_weight_and_old_price():
    html = '''<div class="card"><a href="/product/12"><h3>پیاز زرد ۱ کیلو</h3></a>
      <del>۱۵۰,۰۰۰ تومان</del><span>۱۲۰,۰۰۰</span> <span>تومان</span></div>
      <div class="card"><a href="/about">درباره ما</a><span>۵۰,۰۰۰ تومان</span></div>'''
    found = ps.candidates_from_html(html, "https://www.dastavval.com/search?q=x")
    assert found == [{"title": "پیاز زرد ۱ کیلو", "price": 120000,
                      "url": "https://www.dastavval.com/product/12", "weight_text": ""}]


def test_search_generic_tries_urls_and_remembers(monkeypatch):
    calls = []

    def fake_get(url, params=None, **kwargs):
        calls.append(url)
        if url == "https://mobile.bonakbazar.com/search":
            raise ps.requests.HTTPError("404")
        return FakeResp('[{"title": "عدس", "unit_price": 900000, "unit": "کیلوگرم"}]', url)

    monkeypatch.setattr(ps.session, "get", fake_get)
    monkeypatch.setattr(ps, "WORKING_SEARCH", {})
    assert ps.search_bonakbazar("عدس")[0]["price"] == 90000
    assert calls == ["https://mobile.bonakbazar.com/search", "https://bonakbazar.com/search"]
    calls.clear()
    ps.search_bonakbazar("عدس")
    assert calls == ["https://bonakbazar.com/search"]      # آدرس جواب‌داده به خاطر سپرده شد


def test_search_generic_all_fail_raises(monkeypatch):
    def fake_get(url, params=None, **kwargs):
        raise ps.requests.ConnectionError("no route")

    monkeypatch.setattr(ps.session, "get", fake_get)
    monkeypatch.setattr(ps, "WORKING_SEARCH", {})
    with pytest.raises(ps.requests.ConnectionError):
        ps.search_dastavval("شکر")


def test_comparison_sheet_rial_per_gram(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    monkeypatch.setattr(ps, "DELAY", 0)
    monkeypatch.setattr(ps, "SITES", {
        "parsehshop.ir": lambda q: [{"title": "عدس ۱ کیلویی", "price": 90000, "url": "u1"}],
        "bonakbazar.com": lambda q: [{"title": "عدس درشت", "price": 80000, "url": "u2", "weight_text": "1 کیلوگرم"}],
        "dastavval.com": lambda q: [],
    })
    rows = ps.collect(["عدس"])
    assert [r["rial_per_gram"] for r in rows] == [900.0, 800.0, None]   # ۹۰٬۰۰۰ تومان/کیلو = ۹۰۰ ریال/گرم
    path = tmp_path / "out.xlsx"
    ps.save_excel(rows, path)
    ws = load_workbook(path)["مقایسه (ریال به ازای هر گرم)"]
    assert [c.value for c in ws[1]] == ["نام ماده", "parsehshop.ir", "bonakbazar.com", "dastavval.com", "ارزان‌ترین سایت"]
    assert [c.value for c in ws[2]] == ["عدس", 900, 800, "—", "bonakbazar.com"]
    assert ws["C2"].fill.fgColor.rgb.endswith("C6EFCE") and ws["B2"].hyperlink.target == "u1"
